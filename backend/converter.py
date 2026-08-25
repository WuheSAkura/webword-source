"""Word 公文格式转换引擎 —— 拆分 + 多信号识别（带置信度）+ 按校正结果套用。

流程：
  build_structure(path)            识别每段类型 + 置信度，供人工校正
  convert_with_roles(in,out,roles) 按用户校正后的类型套格式，返回转换日志
所有格式参数来自基础配置与文种模板配置，不在本文件硬编码。
"""

import re
from copy import deepcopy
from collections.abc import Mapping
from pathlib import Path
from typing import Any
from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Pt, Cm
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from config import (
    get_document_config,
    get_template_classification,
    normalize_template_id,
)

ALIGN_MAP = {"left": 0, "center": 1, "right": 2, "justify": 3}
CHAR_WIDTH_CM = 0.565  # 三号字单字宽近似，仅用于无 *Chars 属性时的兜底

# 低置信度阈值：低于此值前端高亮提醒人工复核
LOW_CONFIDENCE = 0.7

# ── 识别正则 ──
RE_H1 = re.compile(r"^[一二三四五六七八九十百]+\s*[、.．]")
RE_H2 = re.compile(r"^[（(]\s*[一二三四五六七八九十百]+\s*[)）]")
RE_H4 = re.compile(r"^[（(]\s*\d+\s*[)）]")
RE_H3 = re.compile(r"^\d+\s*[.．、](?=\s*\S)")
RE_ATTACH = re.compile(r"^附件\s*\d*\s*[:：]?")
RE_RECIPIENT = re.compile(r"[：:]\s*$")  # 主送机关（发文对象）：以冒号结尾的称呼行
RE_DATE = re.compile(r"\d{4}\s*年\s*\d{0,2}\s*月\s*\d{0,2}\s*日")
RE_DATE_FULL = re.compile(r"^\s*\d{4}\s*年\s*\d{0,2}\s*月\s*\d{0,2}\s*日\s*$")
RE_SECURITY = re.compile(r"^(绝密|机密|秘密|内部)(★|▲)?")
RE_YEAR_HEAD = re.compile(r"^\d{4}\s*年")  # 排除“2026年…”被误判为三级标题
RE_FULL_BRACKET = re.compile(r"^[（(].+[)）]$")
RE_DOC_NUMBER = re.compile(r"〔\s*\d{4}\s*〕.*号\s*$")
RE_CARRIER_META = re.compile(r"(签发人|等级[：:]|发电时间|承办单位|抄送[：:]|主送[：:]|内部传真电报|印发\s*$)")
RE_PUNCT_END = re.compile(r"[。！？!?；;，,、：:]\s*$")
RE_TITLE_NOISE = re.compile(r"(签发人|会签|核稿|主办|抄送|印发|电话|联系人|附件[:：]|第\d+页)")
RE_ORG_NAME = re.compile(r"(厅|局|委|办|处|科|院|司|队|站|中心|办公室|委员会|人民政府|公安|法院|检察院|公司|集团)")
RE_SIGNATURE_PREFIX = re.compile(r"^(附件|抄送|印发|联系人|电话|邮编|地址|主送|抄报)")
RE_INLINE_BODY_START = re.compile(
    r"(坚持|深入|全面|认真|切实|持续|加快|推进|加强|严格|围绕|聚焦|"
    r"各地|各单位|要把|要切实|应当|需要|现将|为了|为进一步|按照|根据|经)"
)

HIERARCHY_PATTERNS = {
    "h1": RE_H1,
    "h2": RE_H2,
    "h3": RE_H3,
    "h4": RE_H4,
}

# 文种关键词（用于标题启发式）
TITLE_END = ("通知", "通报", "报告", "请示", "方案", "意见", "决定", "函", "规定",
             "办法", "条例", "细则", "纪要", "批复", "公告", "通告", "命令", "决议",
             "汇报", "计划", "总结", "说明", "情况", "意见稿")
RE_TITLE_KW = re.compile(
    r"(关于|关于印发|转发).{0,80}?"
    r"(通知|通报|报告|请示|方案|意见|决定|函|规定|办法|条例|细则|纪要|批复|公告|通告|命令|决议|汇报|计划|总结|说明|情况)"
)
SUBTITLE_HINT = ("单位", "职务", "部", "局", "厅", "委", "办", "处", "科", "院", "司",
                 "组", "室", "中心", "主任", "局长", "部长", "处长", "科长", "书记")
QUOTE_CHARS = "“”‘’\"'《》〈〉"


def _rule_list(rules: dict, key: str, default: tuple | list = ()) -> tuple[str, ...]:
    values = rules.get(key, default)
    if isinstance(values, str):
        values = [values]
    return tuple(str(item) for item in (values or ()) if str(item))


def _rule_int(rules: dict, key: str, default: int) -> int:
    try:
        return int(rules.get(key, default))
    except (TypeError, ValueError):
        return default


def _ends_with_any(text: str, keywords: tuple[str, ...]) -> bool:
    return any(text.endswith(keyword) or f"的{keyword}" in text[-24:] for keyword in keywords)


def _looks_like_title(text: str, keywords: tuple[str, ...] = ()) -> bool:
    if not (2 <= len(text) <= 110):
        return False
    if RE_TITLE_NOISE.search(text) or RE_DOC_NUMBER.search(text):
        return False
    if keywords and _ends_with_any(text, keywords):
        return True
    if RE_TITLE_KW.search(text):
        return True
    return len(text) <= 40 and any(text.endswith(k) for k in TITLE_END)


def _looks_like_recipient(text: str, recipient_policy: str = "optional") -> bool:
    if recipient_policy == "forbidden":
        return False
    if not (2 <= len(text) <= 80) or not RE_RECIPIENT.search(text):
        return False
    if RE_DATE.search(text) or "。" in text or RE_TITLE_NOISE.search(text):
        return False
    if text.endswith(("如下：", "如下:", "要求：", "要求:", "意见：", "意见:")):
        return False
    return bool(RE_ORG_NAME.search(text) or text.endswith(("同志：", "单位：", "各地：", "各市：", "各县：")))


def _looks_like_signature_unit(text: str, max_len: int = 36) -> bool:
    if not (2 <= len(text) <= max_len):
        return False
    if RE_SIGNATURE_PREFIX.match(text) or RE_PUNCT_END.search(text) or RE_DATE.search(text):
        return False
    return bool(RE_ORG_NAME.search(text))

# ── 段落拆分（沿用：把混排编号拆成逻辑段落）──
RE_SPLIT_MARKER = re.compile(
    r'(?P<h1>[一二三四五六七八九十]{1,3}[、，,])'
    r'|(?P<h2>[（(][一二三四五六七八九十]{1,3}[）)])'
    r'|(?P<h4>[（(]\d{1,2}[）)])'
    r'|(?P<h3>(?<!\d)\d{1,2}[.．、](?=\s*\S))'
    r'|(?P<attach>附件\s*[：:\d])'
)


def is_valid_marker(text: str, match: re.Match) -> bool:
    """过滤示例/引文里的编号，避免把说明性文字拆碎。"""
    start, end = match.start(), match.end()
    prev = text[start - 1] if start > 0 else ""
    next_char = text[end] if end < len(text) else ""
    if start > 0 and not (prev.isspace() or prev in "\n\r。！？；;：:"):
        return False
    if prev in QUOTE_CHARS or next_char in QUOTE_CHARS:
        return False
    if prev == "第" and match.lastgroup == "h1":
        return False
    if match.lastgroup == "h3":
        if prev.isdigit() or (prev.isascii() and prev.isalpha()) or prev in "+-±.":
            return False
        if next_char.isdigit():
            return False
        if prev in "：:" or "附件" in text[max(0, start - 4):start]:
            return False
    if match.lastgroup == "attach" and prev in (QUOTE_CHARS + "如见"):
        return False
    return True


def split_paragraph_text(text: str) -> list[str]:
    """将一个段落文本按编号分隔符拆分为多个逻辑段落"""
    text = re.sub(r"\s*\n+\s*", "\n", text.strip())
    if not text:
        return [text]
    # 整段被括号包裹（副标题、落款联系人等）不拆分，
    # 避免内部数字/顿号（如“张三，”“电话xxx”）被误当作编号切断
    if RE_FULL_BRACKET.match(text):
        return [text]
    markers = [m for m in RE_SPLIT_MARKER.finditer(text) if is_valid_marker(text, m)]
    if not markers:
        return refine_logical_parts([text])
    parts = []
    first_start = markers[0].start()
    if first_start > 0:
        before = text[:first_start].strip()
        if before:
            parts.append(before)
    for idx, marker in enumerate(markers):
        start = marker.start()
        next_start = markers[idx + 1].start() if idx + 1 < len(markers) else len(text)
        segment = text[start:next_start].strip()
        if segment:
            parts.append(segment)
    return refine_logical_parts(parts)


def split_prelude_title(text: str) -> list[str]:
    """从无换行的文头里拆出主标题，避免文号/主送机关吞掉标题。"""
    if RE_H1.match(text) or RE_H2.match(text) or RE_H3.match(text):
        return [text]
    match = RE_TITLE_KW.search(text)
    if not match:
        return [text]
    title = match.group(0).strip()
    if not (4 <= len(title) <= 90):
        return [text]
    before = text[:match.start()].strip()
    after = text[match.end():].strip()
    result = []
    if before:
        result.extend([p for p in re.split(r"\s{2,}|\n+", before) if p.strip()])
    result.append(title)
    if after:
        result.append(after)
    return result


def refine_logical_parts(parts: list[str]) -> list[str]:
    refined = []
    for part in parts:
        for prelude_part in split_prelude_title(part.strip()):
            refined.append(prelude_part)
    return [p for p in refined if p.strip()]


# ── XML 段落操作 ──
def insert_paragraph_after(para_element):
    new_p = OxmlElement("w:p")
    para_element.addnext(new_p)
    return new_p


def insert_paragraph_before(para_element):
    new_p = OxmlElement("w:p")
    para_element.addprevious(new_p)
    return new_p


def split_document_paragraphs(doc: DocxDocument, split_embedded_markers: bool = True) -> list:
    """遍历所有段落，对含多个编号块的段落进行拆分。返回 (element, text) 列表。"""
    result: list[tuple] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            result.append((para._element, text))
            continue
        parts = split_paragraph_text(text) if split_embedded_markers else [text]
        if len(parts) <= 1:
            result.append((para._element, text))
            continue
        para.text = ""
        if parts[0].strip():
            para.add_run(parts[0])
        result.append((para._element, parts[0]))
        elem = para._element
        for part in parts[1:]:
            if not part.strip():
                continue
            elem = insert_paragraph_after(elem)
            r_elem = OxmlElement("w:r")
            t_elem = OxmlElement("w:t")
            t_elem.text = part.strip()
            t_elem.set(qn("xml:space"), "preserve")
            r_elem.append(t_elem)
            elem.append(r_elem)
            result.append((elem, part.strip()))
    return result


# ── 多信号识别 ──
def _para_format_hint(para) -> dict:
    """提取段落原格式线索：是否居中、最大字号、是否加粗。"""
    centered = para.alignment == 1
    max_size = 0.0
    bold = False
    for run in para.runs:
        try:
            if run.font.size:
                max_size = max(max_size, run.font.size.pt)
        except Exception:
            pass
        if run.font.bold:
            bold = True
    return {"centered": centered, "max_size": max_size, "bold": bold}


def classify_one(text: str, index: int, total: int, fmt: dict | None = None,
                 template_id: str | None = None) -> tuple[str, float]:
    """对单个逻辑段落做基础分类，返回 (role, confidence)。多信号：正则+位置+原格式+长度。"""
    t = text.strip()
    fmt = fmt or {}
    rules = get_template_classification(template_id)
    recipient_policy = rules.get("recipient", "optional")
    title_keywords = _rule_list(rules, "title_keywords")
    title_max_index = _rule_int(rules, "title_max_index", 20)
    recipient_max_index = _rule_int(rules, "recipient_max_index", 8)
    signature_unit_max_len = _rule_int(rules, "signature_unit_max_len", 36)
    if not t:
        return "body", 0.0

    near_head = index <= 2
    near_tail = (total - index) <= 5

    if RE_CARRIER_META.search(t) or RE_DOC_NUMBER.search(t):
        return "other", 0.9

    # 1. 强编号特征（顺序：H2/H4 先于 H3，避免“（1）”被当 H3）
    if RE_H1.match(t):
        if index <= title_max_index and _looks_like_title(t, title_keywords):
            return "title", 0.92
        return "h1", 0.95
    if RE_H2.match(t):
        return "h2", 0.95
    if RE_H4.match(t):
        return "h4", 0.95
    if RE_ATTACH.match(t):
        return "attachment_head", 0.95
    if RE_H3.match(t):
        if RE_YEAR_HEAD.match(t):  # 排除“2026年…”
            return "body", 0.6
        # 三级标题 vs 编号列表正文：过长且以句号结尾 → 低置信，交人工
        if len(t) > 40 and t[-1] in "。.":
            return "body", 0.65
        return "h3", 0.9

    # 2. 涉密标识（可选扩展类型）
    if index <= 1 and RE_SECURITY.match(t) and len(t) <= 30:
        return "security", 0.85

    # 3. 成文日期（落款区）
    if RE_DATE_FULL.match(t) and near_tail:
        return "sign_date", 0.95

    # 4. 标题（文头区）。已选文种时，文种词比通用启发式优先级更高。
    if index <= title_max_index and 2 <= len(t) <= 110 and title_keywords:
        if _ends_with_any(t, title_keywords):
            return "title", 0.98
    if near_head and 4 <= len(t) <= 90:
        if _looks_like_title(t, title_keywords):
            return "title", 0.9
        # 原格式线索：居中 + 大字号 + 加粗 → 强指向标题
        if fmt.get("centered") and fmt.get("max_size", 0) >= 18:
            return "title", 0.75 if fmt.get("bold") else 0.7

    # 5. 副标题：文头整段括号文本
    if (rules.get("subtitle", "optional") != "forbidden"
            and near_head and RE_FULL_BRACKET.match(t) and len(t) <= 60):
        return "subtitle", 0.75

    # 5b. 主送机关（发文对象）：文头区、以冒号结尾的称呼行，顶格不缩进
    #     排除“……如下：”等正文引出句，避免误判
    if index <= recipient_max_index and _looks_like_recipient(t, recipient_policy):
        return "recipient", 0.95 if recipient_policy == "required" else 0.85

    # 5c. 落款·联系人：文末括号行，含“联系人/电话”，首行缩进2字（不与日期对齐）
    if near_tail and RE_FULL_BRACKET.match(t) and ("联系" in t or "电话" in t):
        return "sign_contact", 0.85

    # 6. 落款单位：文末短行、非长句
    if near_tail and _looks_like_signature_unit(t, signature_unit_max_len):
        return "sign_unit", 0.75 if rules.get("signature") == "required" else 0.65

    return "body", 0.8


def refine_structure(items: list[dict], template_id: str | None = None) -> None:
    """二次校正：利用相邻关系修正副标题/落款配对。就地修改 items。"""
    n = len(items)
    rules = get_template_classification(template_id)
    recipient_policy = rules.get("recipient", "optional")
    title_keywords = _rule_list(rules, "title_keywords")
    title_max_index = _rule_int(rules, "title_max_index", 20)
    recipient_after_title_window = _rule_int(rules, "recipient_after_title_window", 6)
    signature_unit_max_len = _rule_int(rules, "signature_unit_max_len", 36)

    if recipient_policy == "forbidden":
        for item in items:
            if item["role"] == "recipient":
                item["role"] = "body"
                item["confidence"] = 0.8

    title_start = 0
    # 文种已知时，以最后一个强文种标题为准；其前面的版头、文号、签发信息原样保留。
    if normalize_template_id(template_id) != "generic":
        candidates = [
            i for i, item in enumerate(items[:title_max_index + 1])
            if title_keywords and _ends_with_any(item["fullText"].strip(), title_keywords)
        ]
        if candidates:
            title_idx = candidates[-1]
            leading_limit = max(0, int(rules.get("title_leading_lines", 2)))
            title_start = title_idx
            for i in range(title_idx - 1, max(-1, title_idx - leading_limit - 1), -1):
                text = items[i]["fullText"].strip()
                if (RE_DOC_NUMBER.search(text) or RE_ATTACH.match(text)
                        or RE_CARRIER_META.search(text) or _looks_like_recipient(text, recipient_policy)):
                    break
                if not _looks_like_title(text, title_keywords) and len(text) > 45:
                    break
                title_start = i
            for i in range(title_start):
                if items[i]["role"] != "security":
                    items[i]["role"] = "other"
                    items[i]["confidence"] = 0.95
            for i in range(title_start, title_idx + 1):
                items[i]["role"] = "title"
                items[i]["confidence"] = max(items[i]["confidence"], 0.9)

    # 需要主送机关的文种，以标题后的首个称呼行作为主送机关；版头存在时不再受绝对索引限制。
    if recipient_policy == "required":
        title_end = max((i for i, item in enumerate(items) if item["role"] == "title"), default=-1)
        for i in range(title_end + 1, min(n, title_end + recipient_after_title_window + 1)):
            text = items[i]["fullText"].strip()
            if _looks_like_recipient(text, recipient_policy):
                items[i]["role"] = "recipient"
                items[i]["confidence"] = 0.95
                break

    # 发文对象之前基本都是标题：把误判为正文/编号的段落回归标题
    # （标题可跨多行，但行与行之间不应被识别成正文，避免套格式时夹空行）
    recip_idx = next((i for i, it in enumerate(items) if it["role"] == "recipient"), None)
    if recip_idx is not None:
        for i in range(title_start, recip_idx):
            text = items[i]["fullText"].strip()
            if items[i]["role"] in ("body", "h1", "h2", "h3", "h4", "other") and _looks_like_title(text, title_keywords):
                items[i]["role"] = "title"
                items[i]["confidence"] = max(items[i]["confidence"], 0.9)
    # 副标题必须紧跟标题之后
    for i in range(1, n):
        if items[i]["role"] == "subtitle" and items[i - 1]["role"] != "title":
            # 上一段不是标题：若本段不像括号副标题则回退正文
            items[i]["role"] = "body"
            items[i]["confidence"] = 0.6
    # 多个附件并列：紧跟“附件头”之后、形如“2.×××”的编号行归为“其他附件”
    # （可链式：其他附件后面的编号行仍是其他附件）
    for i in range(1, n):
        if items[i - 1]["role"] not in ("attachment_head", "attachment_other"):
            continue
        t = items[i]["fullText"].strip()
        if (items[i]["role"] in ("h3", "body") and (RE_H3.match(t) or RE_H4.match(t))
                and t[-1:] not in "。.！？!?；;"):
            items[i]["role"] = "attachment_other"
            items[i]["confidence"] = max(items[i]["confidence"], 0.7)

    # 附件之后、落款之前的短组织名不应被附件链吞掉。
    for i in range(1, n):
        if items[i]["role"] == "attachment_other" and _looks_like_signature_unit(
                items[i]["fullText"].strip(), signature_unit_max_len):
            items[i]["role"] = "sign_unit"
            items[i]["confidence"] = max(items[i]["confidence"], 0.72)

    # 落款单位应紧邻成文日期（其上一段）
    date_idx = [i for i, it in enumerate(items) if it["role"] == "sign_date"]
    if date_idx:
        di = date_idx[0]
        if di - 1 >= 0 and items[di - 1]["role"] in ("body", "other", "attachment_other"):
            prev = items[di - 1]
            if _looks_like_signature_unit(prev["fullText"].strip(), signature_unit_max_len):
                prev["role"] = "sign_unit"
                prev["confidence"] = max(prev["confidence"], 0.78)


def build_structure(file_path: str, template_id: str | None = None,
                    split_embedded_markers: bool = True) -> list[dict]:
    """识别文档结构，返回带置信度的清单（供人工校正与套格式共用）。"""
    template_id = normalize_template_id(template_id)
    doc = Document(str(file_path))
    para_items = split_document_paragraphs(doc, split_embedded_markers)

    # 建立 element → Paragraph 映射以读原格式
    para_by_elem = {p._element: p for p in doc.paragraphs}

    logical = [(elem, txt) for elem, txt in para_items if txt.strip()]
    total = len(logical)
    items = []
    for i, (elem, txt) in enumerate(logical):
        para = para_by_elem.get(elem)
        fmt = _para_format_hint(para) if para is not None else {}
        role, conf = classify_one(txt, i, total, fmt, template_id)
        items.append({
            "index": i,
            "role": role,
            "confidence": round(conf, 2),
            "lowConfidence": conf < LOW_CONFIDENCE,
            "text": txt[:40] + ("…" if len(txt) > 40 else ""),
            "fullText": txt,
        })
    refine_structure(items, template_id)
    for it in items:
        it["lowConfidence"] = it["confidence"] < LOW_CONFIDENCE
    return items


# ── 格式应用 ──
def nuke_run_format(rPr):
    if rPr is None:
        return
    tags = ("w:rFonts", "w:sz", "w:szCs", "w:b", "w:bCs", "w:i", "w:iCs",
            "w:u", "w:color", "w:spacing", "w:kern", "w:position", "w:vertAlign",
            "w:strike", "w:dstrike", "w:shadow", "w:outline", "w:emboss",
            "w:imprint", "w:highlight", "w:lang", "w:fitText", "w:effect",
            "w:w", "w:snapToGrid", "w:shd", "w:smallCaps", "w:caps", "w:vanish",
            "w:webHidden", "w:specVanish", "w:em", "w:cs")
    for tag in tags:
        e = rPr.find(qn(tag))
        if e is not None:
            rPr.remove(e)


def set_font(run, east: str, west: str, size: float, bold: bool = False):
    """一个 run 同时设 ascii/hAnsi/cs=西文、eastAsia=中文，Word 按字符自动选字体。"""
    rPr = run._element.get_or_add_rPr()
    nuke_run_format(rPr)
    run.font.name = west
    run.font.size = Pt(size)
    run.font.bold = bold
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:ascii"), west)
    rFonts.set(qn("w:hAnsi"), west)
    rFonts.set(qn("w:cs"), west)


def _set_ind_attr(paragraph, attr: str, value: str):
    # 用 get_or_add_ind 保证 w:ind 在 pPr 中的 schema 顺序正确
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.get_or_add_ind()
    ind.set(qn(attr), value)


def set_first_line_chars(paragraph, chars: int):
    if chars and chars > 0:
        _set_ind_attr(paragraph, "w:firstLineChars", str(int(chars * 100)))


def set_right_chars(paragraph, chars: float):
    if chars and chars > 0:
        _set_ind_attr(paragraph, "w:rightChars", str(int(round(chars * 100))))


def set_left_chars(paragraph, chars: float):
    if chars and chars > 0:
        _set_ind_attr(paragraph, "w:leftChars", str(int(round(chars * 100))))


def style_for(role: str, config: dict) -> dict:
    """从配置取出某 role 的内部样式字典。"""
    styles = config["styles"]
    s = styles.get(role, styles["body"])
    prefix = "title" if role == "title" else "body"
    line = s.get("line_pt", config["spacing"].get(f"{prefix}_line_pt", 32 if prefix == "title" else 28))
    line_rule = s.get("line_rule", config["spacing"].get(f"{prefix}_line_rule", "exact"))
    line_multiple = s.get("line_multiple", config["spacing"].get(f"{prefix}_line_multiple", 1.0))
    return {
        "role": role,
        "font_east": s["cn"],
        "font_west": config["fonts"]["en"],
        "size": s["size_pt"],
        "bold": s.get("bold", False),
        "align": s.get("align", "justify"),
        "line_spacing": line,
        "line_rule": line_rule,
        "line_multiple": line_multiple,
        "first_line_chars": s.get("first_line_chars", 0),
    }


def apply_paragraph_layout(paragraph, style: dict, right_chars_override: float | None = None,
                           left_chars_override: float | None = None):
    align = style.get("align", "justify")
    fmt = paragraph.paragraph_format
    line_rule = style.get("line_rule", "exact")
    if line_rule == "single":
        fmt.line_spacing = 1.0
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    elif line_rule == "multiple":
        fmt.line_spacing = float(style.get("line_multiple", 1.0))
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    else:
        fmt.line_spacing = Pt(style.get("line_spacing", 28))
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.left_indent = None
    fmt.right_indent = None
    fmt.first_line_indent = None
    # 清掉旧的 ind，避免残留
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is not None:
        old_ind = pPr.find(qn("w:ind"))
        if old_ind is not None:
            pPr.remove(old_ind)

    if align == "center_to_date":
        paragraph.alignment = ALIGN_MAP["right"]
        if right_chars_override is not None:
            set_right_chars(paragraph, right_chars_override)
    elif align == "right_indent_4chars":
        paragraph.alignment = ALIGN_MAP["right"]
        set_right_chars(paragraph, 4)
    elif align == "right_indent_2chars":
        # 落款单位/日期：右对齐，离右边距 2 个汉字（右端对齐，不用空格）
        paragraph.alignment = ALIGN_MAP["right"]
        set_right_chars(paragraph, 2)
    else:
        paragraph.alignment = ALIGN_MAP.get(align, 3)
        set_first_line_chars(paragraph, style.get("first_line_chars", 0))
        # 多个附件时，第二个起按"附件："宽度做悬挂缩进，使序号上下对齐
        if left_chars_override is not None:
            set_left_chars(paragraph, left_chars_override)


def clear_runs(paragraph):
    """清除段落内容，覆盖普通 run、超链接及修订记录等文本容器。"""
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def clear_paragraph_run_format(paragraph):
    pPr = paragraph._element.find(qn("w:pPr"))
    if pPr is not None:
        p_rPr = pPr.find(qn("w:rPr"))
        if p_rPr is not None:
            nuke_run_format(p_rPr)


def run_font_signature(run) -> tuple[str, float | None, bool]:
    """提取足以判断标题/正文边界的显式字体特征。"""
    east = ""
    rPr = run._element.rPr
    if rPr is not None:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is not None:
            east = rFonts.get(qn("w:eastAsia")) or ""
    name = east or run.font.name or ""
    size = run.font.size.pt if run.font.size else None
    return name, size, bool(run.font.bold)


def has_strong_style_change(before: tuple[str, float | None, bool],
                            after: tuple[str, float | None, bool]) -> bool:
    before_name, before_size, before_bold = before
    after_name, after_size, after_bold = after
    if before_bold != after_bold:
        return True
    if before_name and after_name and before_name != after_name:
        return True
    return bool(before_size and after_size and abs(before_size - after_size) >= 0.5)


def source_style_boundary(paragraph, text: str, marker_end: int) -> int | None:
    """从原文显式 run 格式切换中寻找标题短语结束位置。"""
    runs = [run for run in paragraph.runs if run.text]
    run_text = "".join(run.text for run in runs)
    if not runs or run_text.strip() != text.strip():
        return None

    leading = len(run_text) - len(run_text.lstrip())
    initial_signature = None
    raw_offset = 0
    for run in runs:
        run_start = raw_offset
        raw_offset += len(run.text)
        if not run.text.strip():
            continue
        signature = run_font_signature(run)
        if initial_signature is None:
            initial_signature = signature
            continue
        boundary = run_start - leading
        if boundary < marker_end + 1 or boundary >= len(text):
            continue
        if has_strong_style_change(initial_signature, signature) and text[boundary:].strip():
            return boundary
    return None


def inline_heading_boundary(paragraph, text: str, role: str, config: dict) -> int | None:
    """确定同段标题短语边界；返回 None 表示整段均为标题。"""
    pattern = HIERARCHY_PATTERNS.get(role)
    marker = pattern.match(text) if pattern else None
    if marker is None:
        return None

    rules = config.get("inline_hierarchy") or {}
    max_heading_chars = rules.get("max_heading_chars") or {}
    max_chars = int(max_heading_chars.get(role, 18 if role in ("h1", "h2") else 14))
    if rules.get("prefer_source_style_boundary", True):
        boundary = source_style_boundary(paragraph, text, marker.end())
        if boundary is not None:
            return boundary

    terminators = str(rules.get("terminators") or "。！？；：.!?;:")
    first_terminator: int | None = None
    for index in range(marker.end(), len(text)):
        if text[index] in terminators and text[index + 1:].strip():
            first_terminator = index + 1
            break

    if first_terminator is not None and first_terminator - marker.end() <= max_chars:
        return first_terminator

    search_end = min(len(text), marker.end() + max_chars + 8)
    soft_area = text[marker.end():search_end]
    for match in RE_INLINE_BODY_START.finditer(soft_area):
        boundary = marker.end() + match.start()
        heading_text = text[:boundary].strip()
        body_text = text[boundary:].strip()
        heading_len = len(heading_text) - marker.end()
        if 2 <= heading_len <= max_chars and body_text:
            return boundary

    if first_terminator is not None:
        return first_terminator
    return None


def clean_text(text: str) -> str:
    text = text.replace("　", "")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"^([一二三四五六七八九十百]+[、.．])\s+", r"\1", text)
    text = re.sub(r"^([（(][一二三四五六七八九十百]+[）)])\s+", r"\1", text)
    text = re.sub(r"^([（(]\d+[）)])\s+", r"\1", text)
    text = re.sub(r"^(\d+[.．、])\s+", r"\1", text)
    return text.strip()


def append_styled_segment(segments: list[tuple[str, dict]], text: str, style: dict):
    if not text:
        return
    if segments and segments[-1][1] == style:
        previous_text, previous_style = segments[-1]
        segments[-1] = (previous_text + text, previous_style)
    else:
        segments.append((text, style))


def _capture_preserved_inline_spans(paragraph, role: str, config: dict) -> list[tuple[int, int, dict]]:
    rules = config.get("preserve_inline") or {}
    enabled = rules.get("enabled", True)
    roles = set(rules.get("roles") or ["body", "recipient", "attachment_head", "attachment_other", "sign_contact"])
    properties = set(rules.get("properties") or ["bold", "italic", "underline", "color", "highlight"])
    if not enabled or role not in roles:
        return []

    spans: list[tuple[int, int, dict]] = []
    offset = 0
    for run in paragraph.runs:
        text = run.text or ""
        start, end = offset, offset + len(text)
        offset = end
        if not text:
            continue
        props = {}
        if "bold" in properties and run.font.bold:
            props["bold"] = True
        if "italic" in properties and run.font.italic:
            props["italic"] = True
        if "underline" in properties and run.font.underline:
            props["underline"] = run.font.underline
        if "color" in properties and run.font.color and run.font.color.rgb:
            props["color"] = run.font.color.rgb
        if "highlight" in properties and run.font.highlight_color:
            props["highlight"] = run.font.highlight_color
        if props:
            spans.append((start, end, props))
    return spans


def _merged_preserved_props(spans: list[tuple[int, int, dict]], start: int, end: int) -> dict:
    props = {}
    for span_start, span_end, span_props in spans:
        if span_start < end and span_end > start:
            props.update(span_props)
    return props


def _apply_preserved_props(run, props: dict):
    if not props:
        return
    if props.get("bold"):
        run.font.bold = True
    if props.get("italic"):
        run.font.italic = True
    if props.get("underline") is not None:
        run.font.underline = props["underline"]
    if props.get("color") is not None:
        run.font.color.rgb = props["color"]
    if props.get("highlight") is not None:
        run.font.highlight_color = props["highlight"]


def configured_body_markers(config: dict) -> tuple[str, ...]:
    """返回按长度倒序排列的正文分项标识。"""
    rules = config.get("inline_hierarchy") or {}
    return tuple(sorted(
        (str(marker) for marker in (rules.get("body_markers") or []) if str(marker)),
        key=len,
        reverse=True,
    ))


def normalize_role_for_text(role: str, text: str, config: dict) -> str:
    """“一是、二是”等是正文分项，不因外部角色误判而套用标题字体。"""
    if role not in HIERARCHY_PATTERNS:
        return role
    normalized = text.lstrip()
    if any(normalized.startswith(marker) for marker in configured_body_markers(config)):
        return "body"
    return role


def body_marker_segments(text: str, body_style: dict, config: dict) -> list[tuple[str, dict]]:
    """仅将“一是、二是”等段内层次标识加粗，正文保持正文样式。"""
    rules = config.get("inline_hierarchy") or {}
    markers = configured_body_markers(config)
    if not markers:
        return [(text, body_style)] if text else []

    marker_style = dict(body_style)
    marker_style["bold"] = bool(rules.get("body_marker_bold", True))
    pattern = re.compile("|".join(re.escape(marker) for marker in markers))
    boundary_chars = "。；！？：:\n"
    segments: list[tuple[str, dict]] = []
    cursor = 0
    for match in pattern.finditer(text):
        previous = match.start() - 1
        while previous >= 0 and text[previous].isspace():
            previous -= 1
        if previous >= 0 and text[previous] not in boundary_chars:
            continue
        append_styled_segment(segments, text[cursor:match.start()], body_style)
        append_styled_segment(segments, match.group(0), marker_style)
        cursor = match.end()
    append_styled_segment(segments, text[cursor:], body_style)
    return segments


def replace_runs_with_segments(paragraph, segments: list[tuple[str, dict]],
                               preserve_spans: list[tuple[int, int, dict]] | None = None):
    clear_runs(paragraph)
    preserve_spans = preserve_spans or []
    cursor = 0
    for text, style in segments:
        if not text:
            continue
        segment_start = cursor
        segment_end = cursor + len(text)
        cursor = segment_end
        cut_points = {segment_start, segment_end}
        for span_start, span_end, _ in preserve_spans:
            if segment_start < span_start < segment_end:
                cut_points.add(span_start)
            if segment_start < span_end < segment_end:
                cut_points.add(span_end)
        ordered = sorted(cut_points)
        for left, right in zip(ordered, ordered[1:]):
            chunk = text[left - segment_start:right - segment_start]
            if not chunk:
                continue
            run = paragraph.add_run(chunk)
            set_font(run, style["font_east"], style["font_west"], style["size"], style["bold"])
            _apply_preserved_props(run, _merged_preserved_props(preserve_spans, left, right))


def apply_role_style(paragraph, text: str, role: str, config: dict,
                     clean_spaces: bool = True, split_inline_heading: bool = True,
                     force_heading_bold: bool = False):
    """统一应用段落布局和段内字体，标题短语与后续正文保持在同一段。"""
    role_style = style_for(role, config)
    if force_heading_bold and role in HIERARCHY_PATTERNS:
        role_style["bold"] = True
    body_style = style_for("body", config)
    apply_paragraph_layout(paragraph, role_style)
    source_text = paragraph.text.strip()
    source_spans = _capture_preserved_inline_spans(paragraph, role, config)
    clear_paragraph_run_format(paragraph)

    segments: list[tuple[str, dict]] = []
    boundary = inline_heading_boundary(paragraph, text, role, config) if split_inline_heading else None
    if role in HIERARCHY_PATTERNS and boundary is not None:
        heading_text = clean_text(text[:boundary]) if clean_spaces else text[:boundary]
        body_text = clean_text(text[boundary:]) if clean_spaces else text[boundary:]
        append_styled_segment(segments, heading_text, role_style)
        segments.extend(body_marker_segments(body_text, body_style, config))
    else:
        normalized = clean_text(text) if clean_spaces else text
        if role == "body":
            segments.extend(body_marker_segments(normalized, body_style, config))
        else:
            append_styled_segment(segments, normalized, role_style)

    output_text = "".join(segment_text for segment_text, _ in segments)
    preserve_spans = source_spans if output_text == source_text else []
    replace_runs_with_segments(paragraph, segments, preserve_spans)


# ── 页面 / 空行 / 页码 ──
def set_page(doc, page: dict):
    margin = page.get("margin_cm", {})
    for sec in doc.sections:
        sec.top_margin = Cm(margin.get("top", 3.7))
        sec.bottom_margin = Cm(margin.get("bottom", 3.5))
        sec.left_margin = Cm(margin.get("left", 2.8))
        sec.right_margin = Cm(margin.get("right", 2.6))
        sec.header_distance = Cm(page.get("header_cm", 0))
        sec.footer_distance = Cm(page.get("footer_cm", 2.5))
        sec.page_width = Cm(page.get("width_cm", 21.0))
        sec.page_height = Cm(page.get("height_cm", 29.7))


def paragraph_element_text(elem) -> str:
    return "".join(t.text or "" for t in elem.findall(".//" + qn("w:t")))


def ensure_blank_paragraphs_before(para_element, count: int):
    existing = 0
    prev = para_element.getprevious()
    while prev is not None and prev.tag == qn("w:p"):
        if paragraph_element_text(prev).strip():
            break
        existing += 1
        prev = prev.getprevious()
    for _ in range(max(count - existing, 0)):
        insert_paragraph_before(para_element)


def remove_blank_paragraphs_before(para_element):
    """删除紧邻 para_element 之前的所有空段落（用于标题各行之间不留空行）。"""
    prev = para_element.getprevious()
    while prev is not None and prev.tag == qn("w:p") and not paragraph_element_text(prev).strip():
        to_remove = prev
        prev = prev.getprevious()
        to_remove.getparent().remove(to_remove)


def _enable_even_odd_headers(doc: DocxDocument):
    settings = doc.settings.element
    if settings.find(qn("w:evenAndOddHeaders")) is None:
        el = OxmlElement("w:evenAndOddHeaders")
        settings.append(el)


def add_page_number_field(paragraph, align_key: str, page_number: dict):
    paragraph.text = ""
    paragraph.alignment = ALIGN_MAP.get(align_key, 1)
    fmt_str = page_number.get("format", "—{PAGE}—")
    prefix, _, suffix = fmt_str.partition("{PAGE}")
    cn = page_number.get("cn", "宋体")
    size = page_number.get("size_pt", 14)
    west = "Times New Roman"

    if prefix:
        r = paragraph.add_run(prefix)
        set_font(r, cn, west, size)

    r = paragraph.add_run()
    fb = OxmlElement("w:fldChar")
    fb.set(qn("w:fldCharType"), "begin")
    r._r.append(fb)
    r = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    r._r.append(instr)
    set_font(r, cn, west, size)
    r = paragraph.add_run()
    fe = OxmlElement("w:fldChar")
    fe.set(qn("w:fldCharType"), "end")
    r._r.append(fe)

    if suffix:
        r = paragraph.add_run(suffix)
        set_font(r, cn, west, size)


def set_page_numbers(doc: DocxDocument, page_number: dict):
    odd_right = page_number.get("odd_right_even_left", True)
    if odd_right:
        _enable_even_odd_headers(doc)
    for sec in doc.sections:
        footer = sec.footer
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        add_page_number_field(para, "right" if odd_right else "center", page_number)
        if odd_right:
            even = sec.even_page_footer
            epara = even.paragraphs[0] if even.paragraphs else even.add_paragraph()
            add_page_number_field(epara, "left", page_number)


# ── 主流程：按校正后的角色套格式 ──
def _display_width(text: str) -> float:
    """粗略显示宽度：中文记 1，数字/字母/标点记 0.5。用于落款居中计算。"""
    w = 0.0
    for ch in text:
        w += 0.5 if ch.isascii() else 1.0
    return w


def convert_with_roles(input_path: str, output_path: str,
                       roles: Mapping[Any, str] | None = None,
                       clean_spaces: bool = True,
                       template_id: str | None = None,
                       split_embedded_markers: bool = True,
                       split_inline_headings: bool = True,
                       force_heading_bold: bool = False) -> dict:
    """按 roles（{逻辑段索引: 角色}）套格式；缺失的段落回退自动识别。返回转换日志。"""
    template_id = normalize_template_id(template_id)
    config = get_document_config(template_id)
    roles = {int(k): v for k, v in (roles or {}).items()}

    doc = Document(str(input_path))
    para_items = split_document_paragraphs(doc, split_embedded_markers)
    para_by_elem = {p._element: p for p in doc.paragraphs}

    logical = [(elem, txt) for elem, txt in para_items if txt.strip()]
    total = len(logical)

    # 先算好每段最终角色（含自动识别兜底），便于跨段落计算（落款居中）
    auto_items = []
    for i, (elem, txt) in enumerate(logical):
        para = para_by_elem.get(elem)
        fmt = _para_format_hint(para) if para is not None else {}
        auto_role, confidence = classify_one(txt, i, total, fmt, template_id)
        auto_items.append({
            "index": i, "role": auto_role, "confidence": confidence,
            "fullText": txt,
        })
    refine_structure(auto_items, template_id)

    final_roles = []
    for i, item in enumerate(auto_items):
        auto_role = item["role"]
        role = roles.get(i, auto_role)
        if role != "other" and role not in config["styles"]:
            role = "body"
        role = normalize_role_for_text(role, item["fullText"], config)
        final_roles.append(role)

    blank = config["blank_lines"]
    seen_title = False
    seen_attachment = False
    seen_signature = False
    log = []

    for i, (elem, txt) in enumerate(logical):
        para = para_by_elem.get(elem)
        if para is None:
            continue
        role = final_roles[i]
        if role == "other":
            # 版头元数据等仍统一页边距与正文字体，避免与正文混排观感割裂
            apply_role_style(
                para, txt, "body", config, clean_spaces,
                split_inline_heading=False,
                force_heading_bold=False,
            )
            log.append({"index": i, "role": role,
                        "corrected": i in roles and roles[i] != auto_items[i]["role"],
                        "text": txt[:30]})
            continue
        # 空行规范化
        if role == "title":
            if not seen_title:
                # 只在标题首行上方留空行
                ensure_blank_paragraphs_before(elem, blank.get("before_title", 2))
                seen_title = True
            else:
                # 标题其余各行：删掉行间空行，只换行不空行
                remove_blank_paragraphs_before(elem)
        elif role == "attachment_head" and not seen_attachment:
            # 附件区与正文之间空行，只在第一段“附件头”前加
            ensure_blank_paragraphs_before(elem, blank.get("before_attachment", 1))
            seen_attachment = True
        elif role == "sign_unit" and not seen_signature:
            ensure_blank_paragraphs_before(elem, blank.get("before_signature", 3))
            seen_signature = True

        apply_role_style(
            para, txt, role, config, clean_spaces,
            split_inline_heading=split_inline_headings,
            force_heading_bold=force_heading_bold,
        )
        log.append({"index": i, "role": role,
                    "corrected": i in roles and roles[i] != auto_items[i]["role"],
                    "text": txt[:30]})

    set_page(doc, config["page"])
    set_page_numbers(doc, config["page_number"])

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return {"output": output_path, "log": log, "template_id": template_id}


def convert_docx(input_path: str, output_path: str, clean_spaces: bool = True,
                 template_id: str | None = None) -> str:
    """无人工校正的一键转换（自动识别 + 套格式），用于兼容/兜底。"""
    result = convert_with_roles(
        input_path, output_path, roles=None,
        clean_spaces=clean_spaces, template_id=template_id,
    )
    return result["output"]


# ── 局部字体微调（套格式后的补充编辑）──
def apply_inline_font(run, font: dict, fallback: dict):
    east = font.get("font_east") or fallback["font_east"]
    west = font.get("font_west") or fallback["font_west"]
    size = font.get("size") or fallback["size"]
    bold = font.get("bold")
    if bold is None:
        bold = fallback["bold"]
    set_font(run, east, west, size, bold)
    if font.get("italic") is not None:
        run.font.italic = bool(font.get("italic"))
    if font.get("underline") is not None:
        run.font.underline = bool(font.get("underline"))


def build_style(role: str | None, overrides: dict | None = None,
                template_id: str | None = None) -> dict:
    config = get_document_config(template_id)
    base_role = role if role is not None and role in config["styles"] else "body"
    style = style_for(base_role, config)
    overrides = overrides or {}
    if overrides.get("alignment"):
        style["align"] = overrides["alignment"]
    if overrides.get("line_spacing") is not None:
        style["line_spacing"] = overrides["line_spacing"]
    if overrides.get("line_rule"):
        style["line_rule"] = overrides["line_rule"]
    elif overrides.get("line_spacing") is not None:
        style["line_rule"] = "exact"
    if overrides.get("line_multiple") is not None:
        style["line_multiple"] = overrides["line_multiple"]
    if overrides.get("first_indent") is not None:
        style["first_line_chars"] = overrides["first_indent"]
    return style


def add_run_with_format(paragraph, text: str, rPr):
    run = paragraph.add_run(text)
    if rPr is not None:
        current = run._element.rPr
        if current is not None:
            run._element.remove(current)
        run._element.insert(0, deepcopy(rPr))
    return run


def rewrite_paragraph_with_selection(paragraph, start, end, base_style, font):
    """只重写选中区间，未选中的 run 格式按原样复制。"""
    text = paragraph.text
    start = max(0, min(start, len(text)))
    end = max(start, min(end, len(text)))

    snapshots = []
    offset = 0
    for run in paragraph.runs:
        if not run.text:
            continue
        run_start = offset
        offset += len(run.text)
        snapshots.append((run_start, offset, run.text, deepcopy(run._element.rPr)))

    if not snapshots or "".join(item[2] for item in snapshots) != text:
        snapshots = [(0, len(text), text, None)]

    clear_runs(paragraph)
    for run_start, run_end, run_text, rPr in snapshots:
        cut_points = {run_start, run_end}
        if run_start < start < run_end:
            cut_points.add(start)
        if run_start < end < run_end:
            cut_points.add(end)
        ordered = sorted(cut_points)
        for left, right in zip(ordered, ordered[1:]):
            segment = run_text[left - run_start:right - run_start]
            if not segment:
                continue
            selected = left < end and right > start
            run = add_run_with_format(paragraph, segment, rPr)
            if selected and font:
                apply_inline_font(run, font, base_style)


def apply_local_edit(docx_path: str, selection: dict, font: dict | None = None,
                     paragraph: dict | None = None,
                     template_id: str | None = None) -> str:
    doc = Document(str(docx_path))
    # 与预览/结构面板一致：仅对非空逻辑段落编号，避免空段导致索引错位
    logical_paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    start_para = int(selection.get("startParagraph", 0))
    end_para = int(selection.get("endParagraph", start_para))
    start_offset = int(selection.get("startOffset", 0))
    end_offset = int(selection.get("endOffset", 0))
    if start_para > end_para:
        start_para, end_para = end_para, start_para
        start_offset, end_offset = end_offset, start_offset
    if not logical_paragraphs:
        raise ValueError("文档没有可编辑段落")
    start_para = max(0, min(start_para, len(logical_paragraphs) - 1))
    end_para = max(0, min(end_para, len(logical_paragraphs) - 1))
    total = len(logical_paragraphs)

    paragraph = paragraph or {}
    role_override = paragraph.get("role") or None

    for para_index in range(start_para, end_para + 1):
        para = logical_paragraphs[para_index]
        text = para.text
        if not text:
            continue
        role = role_override or classify_one(text, para_index, total, template_id=template_id)[0]
        base_style = build_style(role, paragraph, template_id)
        apply_paragraph_layout(para, base_style)
        seg_start = start_offset if para_index == start_para else 0
        seg_end = end_offset if para_index == end_para else len(text)
        if seg_end < seg_start:
            seg_start, seg_end = seg_end, seg_start
        rewrite_paragraph_with_selection(para, seg_start, seg_end, base_style, font)

    doc.save(str(docx_path))
    return docx_path
