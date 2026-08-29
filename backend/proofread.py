"""公文纠错：抽取正文、调用 DeepSeek 识别问题、按字眼替换导出。"""

from __future__ import annotations

import json
import re
import shutil
import uuid
from pathlib import Path
from time import perf_counter
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ai_writer import (
    MODEL_HTTP_TIMEOUT_SECONDS,
    SUPPORTED_INPUT_SUFFIXES,
    UNREADABLE_REFERENCE_SUFFIXES,
    call_chat_model_result,
    get_model_config,
    safe_extract_text,
)

BASE_DIR = Path(__file__).resolve().parent
TEMP_DIR = BASE_DIR / "temp"
TEMP_DIR.mkdir(exist_ok=True)
LIBRARY_PATH = BASE_DIR / "proofread_library.yaml"

PROOFREAD_SUFFIXES = SUPPORTED_INPUT_SUFFIXES | UNREADABLE_REFERENCE_SUFFIXES | {".doc"}
MAX_PROOFREAD_UPLOAD = 30 * 1024 * 1024

# session_id -> session dict
proofread_sessions: dict[str, dict[str, Any]] = {}
_library_cache: dict[str, Any] | None = None

ALIGN_MAP = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    WD_ALIGN_PARAGRAPH.DISTRIBUTE: "justify",
}

BASE_SYSTEM_PROMPT = """你是党政机关公文校对专家。必须先通读所给全部段落、结合上下文理解文意后再纠错。

工作方式（正向）：
1. 先理解文意与指代对象，再对照「公文纠错知识库」中的正确写法与前提条件进行核对；
2. 命中知识库前提时，按库中正确示例与正误对照给出修改；
3. 只报告有效必要的问题；不做语气润色，不改无意义的空格/全角半角偏好，不做敏感政治审查。

核对范围（按优先级，结合语境）：
1. 错别字、多字、漏字、别字
2. 标点与符号的规范用法（含知识库中的间隔号等专项）
3. 数字与计量明显不规范
4. 称谓与落款明显不当
5. 公文套语误用
6. 有上下文依据的前后矛盾、指代不清
7. 日期时间在「成文/落款」等场景下的规范写法

分级：
- suggest（黄）：建议修改（默认）
- error（红）：重大且确定性硬伤（错别字、明显笔误等）；宁缺毋滥，拿不准用 suggest。

输出严格 JSON（不要 markdown 代码围栏；思考过程不要写入 JSON）：
{"issues":[{"paragraphIndex":0,"original":"原文精确片段","suggestion":"建议替换为","level":"error|suggest","comment":"不超过40字说明"}]}

规则：
- original 必须是对应段落中的连续原文子串
- suggestion 只改必要文字
- 同一处只报一条；无问题则 issues 为空数组
- paragraphIndex 从 0 起
"""


def load_proofread_library(force: bool = False) -> dict[str, Any]:
    global _library_cache
    if _library_cache is not None and not force:
        return _library_cache
    data: dict[str, Any] = {"version": 0, "positive_method": "", "rules": []}
    if LIBRARY_PATH.exists():
        try:
            import yaml

            loaded = yaml.safe_load(LIBRARY_PATH.read_text(encoding="utf-8")) or {}
            if isinstance(loaded, dict):
                data = loaded
        except Exception:
            pass
    if not isinstance(data.get("rules"), list):
        data["rules"] = []
    _library_cache = data
    return data


def format_library_for_prompt(library: dict[str, Any] | None = None) -> str:
    """把纠错库收成正向提示：前提 + 正确写法 + 正误对照。"""
    library = library or load_proofread_library()
    parts: list[str] = ["【公文纠错知识库——请按下列正确规范核对】"]
    method = str(library.get("positive_method") or "").strip()
    if method:
        parts.append(method.strip())
    for rule in library.get("rules") or []:
        if not isinstance(rule, dict):
            continue
        title = str(rule.get("title") or rule.get("id") or "条目").strip()
        premise = str(rule.get("premise") or "").strip()
        block = [f"■ {title}"]
        if premise:
            block.append(f"前提：{premise}")
        corrects = [str(x).strip() for x in (rule.get("correct_examples") or []) if str(x).strip()]
        if corrects:
            block.append("正确示例：" + "；".join(f"「{c}」" for c in corrects[:6]))
        pairs = []
        for item in rule.get("incorrect_examples") or []:
            if not isinstance(item, dict):
                continue
            wrong = str(item.get("wrong") or "").strip()
            right = str(item.get("right") or "").strip()
            if wrong and right:
                pairs.append(f"「{wrong}」→「{right}」")
        if pairs:
            block.append("正误对照：" + "；".join(pairs[:6]))
        parts.append("\n".join(block))
    return "\n\n".join(parts)


def build_system_prompt() -> str:
    lib_text = format_library_for_prompt()
    if lib_text.strip():
        return BASE_SYSTEM_PROMPT + "\n\n" + lib_text
    return BASE_SYSTEM_PROMPT


def apply_library_patches(paragraphs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """按知识库确定性补丁扫描正文，生成高置信 issues（可与模型结果合并）。"""
    library = load_proofread_library()
    issues: list[dict[str, Any]] = []
    seen: set[tuple[int, str, str]] = set()
    for rule in library.get("rules") or []:
        if not isinstance(rule, dict):
            continue
        level = str(rule.get("level") or "suggest").strip().lower()
        if level not in {"error", "suggest"}:
            level = "suggest"
        title = str(rule.get("title") or rule.get("id") or "纠错库").strip()
        for patch in rule.get("patches") or []:
            if not isinstance(patch, dict):
                continue
            pattern = str(patch.get("pattern") or "")
            replacement = str(patch.get("replacement") or "")
            if not pattern:
                continue
            try:
                regex = re.compile(pattern)
            except re.error:
                continue
            for para in paragraphs:
                text = str(para.get("text") or "")
                if not text.strip():
                    continue
                try:
                    para_index = int(para.get("index"))
                except (TypeError, ValueError):
                    continue
                for match in regex.finditer(text):
                    original = match.group(0)
                    try:
                        suggestion = match.expand(replacement)
                    except re.error:
                        continue
                    if not original or suggestion == original:
                        continue
                    key = (para_index, original, suggestion)
                    if key in seen:
                        continue
                    seen.add(key)
                    note = ""
                    for item in rule.get("incorrect_examples") or []:
                        if isinstance(item, dict) and str(item.get("wrong") or "").find(original) >= 0:
                            note = str(item.get("note") or "").strip()
                            break
                    comment = note or f"按规范应写作「{suggestion}」（{title}）"
                    if len(comment) > 40:
                        comment = comment[:40]
                    issues.append({
                        "id": uuid.uuid4().hex[:12],
                        "paragraphIndex": para_index,
                        "original": original,
                        "suggestion": suggestion,
                        "level": level,
                        "comment": comment,
                        "applied": False,
                        "source": "library",
                        "ruleId": str(rule.get("id") or ""),
                    })
    return issues


def merge_issues(*groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """合并多来源问题；同一原文片段优先保留纠错库结果。"""
    by_span: dict[tuple[int, str], dict[str, Any]] = {}
    ordered: list[dict[str, Any]] = []
    for group in groups:
        for issue in group:
            if not isinstance(issue, dict):
                continue
            try:
                para_index = int(issue.get("paragraphIndex"))
            except (TypeError, ValueError):
                continue
            original = str(issue.get("original") or "")
            if not original:
                continue
            key = (para_index, original)
            existing = by_span.get(key)
            if existing is None:
                by_span[key] = issue
                ordered.append(issue)
                continue
            # 库规则覆盖模型同跨度结果
            if issue.get("source") == "library" and existing.get("source") != "library":
                idx = ordered.index(existing)
                ordered[idx] = issue
                by_span[key] = issue
    return ordered


def _file_magic(path: Path) -> bytes:
    try:
        with open(path, "rb") as handle:
            return handle.read(8)
    except OSError:
        return b""


def _is_ole_compound(path: Path) -> bool:
    magic = _file_magic(path)
    return magic[:4] == b"\xd0\xcf\x11\xe0"


def _is_zip_docx(path: Path) -> bool:
    magic = _file_magic(path)
    return magic[:2] == b"PK"


def _align_of(paragraph) -> str:
    try:
        align = paragraph.alignment
        if align is None and paragraph.paragraph_format.alignment is not None:
            align = paragraph.paragraph_format.alignment
        return ALIGN_MAP.get(align, "left") if align is not None else "left"
    except Exception:
        return "left"


def _para_item(index: int, text: str, align: str = "left", *, source: str = "paragraph") -> dict[str, Any]:
    value = text if text is not None else ""
    return {
        "index": index,
        "text": value,
        "align": align or "left",
        "empty": not bool(str(value).strip()),
        "source": source,
    }


def extract_docx_paragraphs(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    """抽取 docx 正文：段落 + 表格单元格，尽量保证可预览。"""
    warnings: list[str] = []
    if _is_ole_compound(path):
        raise ValueError(
            "该文件是旧版 Word（.doc / OLE）或未真正另存为 .docx。"
            "请用 Word/WPS「另存为」成 .docx 后再上传。"
        )
    if not _is_zip_docx(path):
        raise ValueError("文件不是有效的 .docx（ZIP）格式，请重新另存为 .docx 后再试。")

    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ValueError(
            f"无法打开 Word 文档（可能加密、损坏或非标准 docx）：{exc}"
        ) from exc

    items: list[dict[str, Any]] = []
    for index, para in enumerate(doc.paragraphs):
        items.append(_para_item(index, para.text or "", _align_of(para), source="paragraph"))

    table_added = 0
    try:
        from ai_writer import _iter_body_blocks

        next_index = len(items)
        for block_type, block_index, block in _iter_body_blocks(doc):
            if block_type != "table":
                continue
            seen_cells: set[Any] = set()
            try:
                rows = list(block.rows)
            except Exception as exc:
                warnings.append(f"表格 {block_index} 读取失败，已跳过：{exc}")
                continue
            for row in rows:
                try:
                    cells = list(row.cells)
                except Exception:
                    continue
                for cell in cells:
                    cell_key = cell._tc
                    if cell_key in seen_cells:
                        continue
                    seen_cells.add(cell_key)
                    cell_text = "\n".join(p.text for p in cell.paragraphs).strip()
                    if not cell_text:
                        continue
                    for line in cell_text.split("\n"):
                        line = line.strip()
                        if not line:
                            continue
                        items.append(_para_item(next_index, line, "left", source="table"))
                        next_index += 1
                        table_added += 1
    except Exception as exc:
        warnings.append(f"表格抽取降级：{exc}")

    if table_added:
        warnings.append(f"已从表格补充 {table_added} 行正文用于预览与纠错")

    non_empty = sum(1 for item in items if not item["empty"])
    if non_empty == 0:
        try:
            from ai_writer import extract_document_inventory, inventory_to_text

            inventory = extract_document_inventory(path)
            text = inventory_to_text(inventory).strip()
            warnings.extend(list(inventory.get("warnings") or []))
            if text:
                fallback = paragraphs_from_plain_text(text)
                for item in fallback:
                    item["source"] = "fallback"
                warnings.append("原文档段落为空，已用兜底方式提取正文预览")
                return fallback, warnings
        except Exception as exc:
            warnings.append(f"兜底抽取失败：{exc}")
        warnings.append("未能抽取到可读正文（可能是扫描件/纯图片文档）")

    return items, warnings


def paragraphs_from_plain_text(text: str) -> list[dict[str, Any]]:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if not lines:
        lines = [""]
    return [_para_item(i, line, "left", source="text") for i, line in enumerate(lines)]


def load_source_paragraphs(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    suffix = path.suffix.lower()
    warnings: list[str] = []
    if suffix == ".docx":
        return extract_docx_paragraphs(path)
    if suffix == ".doc":
        if _is_ole_compound(path):
            raise ValueError("暂不支持旧版 .doc，请先另存为 .docx 后再上传。")
        # 有些文件扩展名是 .doc 实际是 docx
        if _is_zip_docx(path):
            warnings.append("扩展名为 .doc，但内容是 docx，已按 docx 打开")
            return extract_docx_paragraphs(path)
        raise ValueError("无法识别的 .doc 文件，请另存为 .docx 后再试。")
    if suffix == ".txt":
        raw = path.read_text(encoding="utf-8", errors="ignore")
        return paragraphs_from_plain_text(raw), warnings
    if suffix in UNREADABLE_REFERENCE_SUFFIXES:
        warnings.append(f"{suffix} 暂无法直接抽取正文，请先转为 docx/pdf/txt")
        return paragraphs_from_plain_text(""), warnings
    text, extract_warnings = safe_extract_text(path, limit=None)
    warnings.extend(extract_warnings)
    if not text.strip():
        warnings.append("未能抽取到正文，请确认文件可读或先转为 docx")
    return paragraphs_from_plain_text(text), warnings


def set_paragraph_text_keep_style(paragraph, new_text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return
    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


def apply_paragraphs_to_docx(source_path: Path, paragraphs: list[dict[str, Any]], dest_path: Path) -> Path:
    shutil.copy2(source_path, dest_path)
    doc = Document(str(dest_path))
    doc_paras = list(doc.paragraphs)
    pending_replaces: list[tuple[str, str]] = []
    for item in paragraphs:
        try:
            index = int(item.get("index", -1))
        except (TypeError, ValueError):
            continue
        new_text = item.get("text")
        if new_text is None:
            continue
        source = str(item.get("source") or "paragraph")
        if source == "paragraph" and 0 <= index < len(doc_paras):
            if (doc_paras[index].text or "") == new_text:
                continue
            set_paragraph_text_keep_style(doc_paras[index], str(new_text))
            continue
        # 表格/兜底段落：按字眼替换写回全文
        original = str(item.get("originalText") or "")
        if original and original != new_text and original in "".join(p.text for p in doc.paragraphs):
            pending_replaces.append((original, str(new_text)))

    if pending_replaces:
        for para in doc.paragraphs:
            text = para.text or ""
            updated = text
            for old, new in pending_replaces:
                if old and old in updated:
                    updated = updated.replace(old, new)
            if updated != text:
                set_paragraph_text_keep_style(para, updated)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        text = para.text or ""
                        updated = text
                        for old, new in pending_replaces:
                            if old and old in updated:
                                updated = updated.replace(old, new)
                        if updated != text:
                            set_paragraph_text_keep_style(para, updated)

    doc.save(str(dest_path))
    return dest_path


def export_as_docx_from_paragraphs(paragraphs: list[dict[str, Any]], dest_path: Path, title: str = "纠错结果") -> Path:
    doc = Document()
    wrote = False
    for item in paragraphs:
        text = str(item.get("text") or "")
        para = doc.add_paragraph(text)
        align = str(item.get("align") or "left")
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "justify":
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        wrote = True
    if not wrote:
        doc.add_paragraph(title)
    doc.save(str(dest_path))
    return dest_path


def export_as_txt(paragraphs: list[dict[str, Any]], dest_path: Path) -> Path:
    content = "\n".join(str(item.get("text") or "") for item in paragraphs)
    dest_path.write_text(content, encoding="utf-8")
    return dest_path


def _strip_json_fence(raw: str) -> str:
    text = (raw or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def _loads_json_lenient(text: str) -> Any:
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    # 截断 JSON：尝试补全常见收尾
    repaired = text.rstrip()
    if repaired.endswith(","):
        repaired = repaired[:-1]
    for suffix in ('"]}', "]}", "}", "]}", "]}"):
        try:
            return json.loads(repaired + suffix)
        except json.JSONDecodeError:
            continue
    # 提取 issues 数组片段
    match = re.search(r'"issues"\s*:\s*(\[[\s\S]*)', text)
    if match:
        arr = match.group(1)
        # 截到最后一个完整对象
        last_obj = arr.rfind("}")
        if last_obj >= 0:
            candidate = arr[: last_obj + 1] + "]"
            try:
                return {"issues": json.loads(candidate)}
            except json.JSONDecodeError:
                pass
    match = re.search(r"\{[\s\S]*\}", text)
    if match:
        return json.loads(match.group(0))
    raise json.JSONDecodeError("无法解析", text, 0)


def _parse_issues_payload(raw: str) -> tuple[list[dict[str, Any]], str | None]:
    """返回 (issues, warning)。尽量不因模型截断而整体失败。"""
    text = _strip_json_fence(raw)
    if not text:
        return [], "模型返回空内容"
    try:
        data = _loads_json_lenient(text)
    except Exception:
        # 再试：从文本里捞多条 JSON 对象
        objs = re.findall(r"\{[^{}]*\"original\"[^{}]*\}", text)
        issues = []
        for chunk in objs:
            try:
                item = json.loads(chunk)
            except json.JSONDecodeError:
                continue
            if isinstance(item, dict):
                issues.append(item)
        if issues:
            return issues, "模型输出不完整，已尽力解析部分结果"
        return [], "模型未返回可解析的 JSON 纠错结果"
    if isinstance(data, dict):
        issues = data.get("issues")
    else:
        issues = data
    if issues is None:
        return [], "纠错结果缺少 issues 字段"
    if not isinstance(issues, list):
        return [], "纠错结果格式错误：issues 不是数组"
    return [item for item in issues if isinstance(item, dict)], None


def normalize_issues(raw_issues: list[dict[str, Any]], paragraphs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    para_map = {int(p["index"]): str(p.get("text") or "") for p in paragraphs}
    normalized: list[dict[str, Any]] = []
    seen: set[tuple[int, str, str]] = set()
    soft_error_markers = (
        "引号", "语气", "润色", "更好", "建议改为", "可读", "简洁", "口语",
        "格式", "空格", "全角", "半角", "排版", "偏好", "可选",
    )
    for item in raw_issues:
        try:
            para_index = int(item.get("paragraphIndex", item.get("paragraph_index", -1)))
        except (TypeError, ValueError):
            continue
        original = str(item.get("original") or "").strip()
        suggestion = str(item.get("suggestion") or "").strip()
        if para_index not in para_map or not original:
            continue
        if original not in para_map[para_index]:
            continue
        if suggestion == original:
            continue
        level = str(item.get("level") or "suggest").strip().lower()
        if level not in {"error", "suggest"}:
            level = "suggest"
        comment = str(item.get("comment") or item.get("reason") or "").strip()[:80]
        source = str(item.get("source") or "model")
        # 软性理由不允许标红（纠错库条目除外）
        if source != "library" and level == "error" and any(marker in comment for marker in soft_error_markers):
            level = "suggest"
        # 仅改标点/引号且无明确「错别字」类说明 → 降为建议（纠错库间隔号等专项除外）
        if source != "library" and level == "error":
            only_punct = bool(re.fullmatch(
                r"[\s\"'“”‘’《》〈〉（）()【】\[\]、，。；：！？,.!?:;…—\-·]+",
                original + suggestion,
            ))
            if only_punct and ("错别字" not in comment and "别字" not in comment):
                level = "suggest"
        key = (para_index, original, suggestion)
        if key in seen:
            continue
        seen.add(key)
        entry = {
            "id": str(item.get("id") or uuid.uuid4().hex[:12]),
            "paragraphIndex": para_index,
            "original": original,
            "suggestion": suggestion,
            "level": level,
            "comment": comment,
            "applied": False,
            "source": source,
        }
        if item.get("ruleId"):
            entry["ruleId"] = str(item.get("ruleId"))
        normalized.append(entry)
    # 全文红色上限：过多则只保留前若干条，其余降级为建议（纠错库命中除外）
    max_errors = 3
    errors = [x for x in normalized if x["level"] == "error" and x.get("source") != "library"]
    if len(errors) > max_errors:
        for extra in errors[max_errors:]:
            extra["level"] = "suggest"
    # 同一段落最多 1 条红色（纠错库命中除外）
    seen_para_error: set[int] = set()
    for issue in normalized:
        if issue["level"] != "error":
            continue
        if issue.get("source") == "library":
            continue
        para_index = issue["paragraphIndex"]
        if para_index in seen_para_error:
            issue["level"] = "suggest"
        else:
            seen_para_error.add(para_index)
    return normalized


def build_model_user_content(paragraphs: list[dict[str, Any]], *, doc_hint: str = "") -> str:
    lines = [
        "请先通读下列公文段落并理解上下文，再按系统提示与纠错知识库中的正确写法核对。",
        "条款与地名以材料语境为准；命中知识库前提时按库中正确示例修改。",
        "按系统要求只返回 JSON（issues 数组）。",
        "",
    ]
    if doc_hint:
        lines.append(doc_hint)
        lines.append("")
    for item in paragraphs:
        idx = item["index"]
        text = item.get("text") or ""
        if not str(text).strip():
            continue
        lines.append(f"[{idx}] {text}")
    lines.append("")
    lines.append("只输出 JSON。")
    return "\n".join(lines)


def _chunk_paragraphs(paragraphs: list[dict[str, Any]], max_chars: int = 12000) -> list[list[dict[str, Any]]]:
    chunks: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    size = 0
    for item in paragraphs:
        text = str(item.get("text") or "")
        if not text.strip():
            continue
        add = len(text) + 8
        if current and size + add > max_chars:
            chunks.append(current)
            current = []
            size = 0
        current.append(item)
        size += add
    if current:
        chunks.append(current)
    return chunks or [[]]


def _extract_issues_from_model_result(result: dict[str, Any]) -> tuple[list[dict[str, Any]], str | None]:
    """优先解析 content；思考模式下 content 为空或截断时再从 reasoning 捞 JSON。"""
    content = (result.get("content") or "").strip()
    reasoning = (result.get("reasoning") or "").strip()
    raw_issues, warning = _parse_issues_payload(content)
    if raw_issues:
        return raw_issues, warning
    if reasoning:
        alt_issues, alt_warn = _parse_issues_payload(reasoning)
        if alt_issues:
            return alt_issues, alt_warn or "已从思考内容中提取 JSON 结果"
        # 思考文本里常夹带最终 JSON：再扫一次带 issues 的对象
        match = re.search(r"\{[\s\S]*\"issues\"\s*:\s*\[[\s\S]*\}\s*$", reasoning)
        if not match:
            match = re.search(r"\{[\s\S]*\"issues\"\s*:\s*\[[\s\S]*\}", reasoning)
        if match:
            alt_issues, alt_warn = _parse_issues_payload(match.group(0))
            if alt_issues:
                return alt_issues, alt_warn or "已从思考内容中提取 JSON 结果"
    if warning:
        return [], warning
    if not content and not reasoning:
        return [], "模型返回空内容"
    return [], "模型未返回可解析的 JSON 纠错结果"


def run_proofread_with_model(paragraphs: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    started = perf_counter()
    cfg = get_model_config(mask_secret=False)
    request_url = cfg.get("requestUrl") or ""
    api_key = cfg.get("apiKey") or ""
    model_name = cfg.get("modelName") or "deepseek-v4-flash"
    if not (api_key or "").strip() and not (request_url or "").strip():
        raise ValueError("未配置模型服务，请先在「公文写作」中配置 DeepSeek")

    non_empty = [p for p in paragraphs if str(p.get("text") or "").strip()]
    if not non_empty:
        return [], {
            "warning": "文档没有可纠错的正文，请确认预览中有文字",
            "issueCount": 0,
            "errorCount": 0,
            "suggestCount": 0,
            "elapsedSeconds": round(perf_counter() - started, 2),
        }

    chunks = _chunk_paragraphs(non_empty)
    all_raw: list[dict[str, Any]] = []
    warnings: list[str] = []
    last_result: dict[str, Any] = {}
    total_paras = len(non_empty)
    system_prompt = build_system_prompt()
    library_started = perf_counter()
    library_issues = apply_library_patches(paragraphs)
    library_elapsed = round(perf_counter() - library_started, 2)
    chunk_timings: list[dict[str, Any]] = []
    for index, chunk in enumerate(chunks, start=1):
        chunk_started = perf_counter()
        retried = False
        doc_hint = ""
        if len(chunks) > 1:
            doc_hint = (
                f"说明：全文约 {total_paras} 个有效段落，当前为第 {index}/{len(chunks)} 批。"
                "不同段落出现的地名/单位名可能指不同对象，应分别保留；"
                "法律条款以材料表述为准，仅在有把握的笔误时修改。"
            )
        messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": build_model_user_content(chunk, doc_hint=doc_hint)},
        ]
        try:
            # 允许思考：提高 max_tokens，避免 reasoning 占满额度导致 JSON 被截断/为空
            result = call_chat_model_result(
                request_url=request_url,
                api_key=api_key,
                model_name=model_name,
                messages=messages,
                temperature=0.1,
                max_tokens=16384,
                thinking="enabled",
                response_format={"type": "json_object"},
            )
            last_result = result
            raw_issues, parse_warning = _extract_issues_from_model_result(result)
            # 思考占满仍无 JSON 时，该批降级重试（关思考），避免整单失败
            if not raw_issues and (
                result.get("finishReason") == "length"
                or (parse_warning and ("未返回" in parse_warning or "空内容" in parse_warning or "无法解析" in parse_warning))
            ):
                retried = True
                warnings.append(f"第{index}批：思考输出未得到有效 JSON，已无思考重试")
                result = call_chat_model_result(
                    request_url=request_url,
                    api_key=api_key,
                    model_name=model_name,
                    messages=messages,
                    temperature=0.1,
                    max_tokens=4096,
                    thinking="disabled",
                    reasoning_effort="none",
                    response_format={"type": "json_object"},
                )
                last_result = result
                raw_issues, parse_warning = _extract_issues_from_model_result(result)
            if parse_warning and not raw_issues:
                warnings.append(f"第{index}批：{parse_warning}")
            elif parse_warning:
                warnings.append(f"第{index}批：{parse_warning}")
            if result.get("finishReason") == "length" and raw_issues:
                warnings.append(f"第{index}批：输出可能被截断，结果或不全")
            all_raw.extend(raw_issues)
            chunk_timings.append({
                "index": index,
                "paragraphCount": len(chunk),
                "elapsedSeconds": round(perf_counter() - chunk_started, 2),
                "issueCount": len(raw_issues),
                "reasoningChars": int(result.get("reasoningChars") or 0),
                "retried": retried,
                "finishReason": result.get("finishReason"),
            })
        except Exception as exc:
            warnings.append(f"第{index}批纠错失败：{exc}")
            chunk_timings.append({
                "index": index,
                "paragraphCount": len(chunk),
                "elapsedSeconds": round(perf_counter() - chunk_started, 2),
                "issueCount": 0,
                "error": str(exc),
            })
            continue

    model_issues = normalize_issues(all_raw, paragraphs)
    # 模型返回 0 条且各批均成功时视为「无问题」，允许再次纠错；仅全部批次失败才报错
    successful_chunks = [t for t in chunk_timings if not t.get("error")]
    if not model_issues and not library_issues and not successful_chunks:
        raise RuntimeError("；".join(warnings[:3]) if warnings else "纠错未产生结果")

    issues = merge_issues(library_issues, model_issues)
    # 再走一遍定位校验，确保仍落在原文中
    issues = normalize_issues(issues, paragraphs)
    elapsed = round(perf_counter() - started, 2)
    model_elapsed = round(sum(t.get("elapsedSeconds") or 0 for t in chunk_timings), 2)
    meta = {
        "model": last_result.get("model") or model_name,
        "finishReason": last_result.get("finishReason"),
        "usage": last_result.get("usage") or {},
        "issueCount": len(issues),
        "errorCount": sum(1 for i in issues if i["level"] == "error"),
        "suggestCount": sum(1 for i in issues if i["level"] == "suggest"),
        "chunkCount": len(chunks),
        "chunkTimings": chunk_timings,
        "elapsedSeconds": elapsed,
        "libraryElapsedSeconds": library_elapsed,
        "modelElapsedSeconds": model_elapsed,
        "modelHttpTimeoutSeconds": MODEL_HTTP_TIMEOUT_SECONDS,
        "warnings": warnings,
        "reasoningChars": last_result.get("reasoningChars") or 0,
        "libraryHits": sum(1 for i in issues if i.get("source") == "library"),
        "libraryVersion": load_proofread_library().get("version"),
    }
    if warnings:
        meta["warning"] = "；".join(warnings[:4])
    return issues, meta


SESSION_DIR = TEMP_DIR / "proofread_sessions"
SESSION_DIR.mkdir(exist_ok=True)


def _session_path(session_id: str) -> Path:
    return SESSION_DIR / f"{session_id}.json"


def persist_session(session: dict[str, Any]) -> None:
    try:
        payload = {
            "id": session["id"],
            "sourcePath": session.get("sourcePath"),
            "sourceName": session.get("sourceName"),
            "sourceKind": session.get("sourceKind"),
            "sourceSuffix": session.get("sourceSuffix"),
            "platformFileId": session.get("platformFileId"),
            "paragraphs": session.get("paragraphs") or [],
            "issues": session.get("issues") or [],
            "status": session.get("status") or "ready",
            "warnings": session.get("warnings") or [],
            "meta": session.get("meta") or {},
        }
        _session_path(session["id"]).write_text(
            json.dumps(payload, ensure_ascii=False),
            encoding="utf-8",
        )
    except OSError:
        pass


def _load_session_from_disk(session_id: str) -> dict[str, Any] | None:
    path = _session_path(session_id)
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if not isinstance(data, dict) or not data.get("id"):
        return None
    proofread_sessions[session_id] = data
    return data


def create_session(
    *,
    source_path: Path,
    source_name: str,
    source_kind: str,
    platform_file_id: str | None = None,
) -> dict[str, Any]:
    paragraphs, warnings = load_source_paragraphs(source_path)
    # 记录原始文本，便于表格类段落按字眼回写
    for item in paragraphs:
        item["originalText"] = str(item.get("text") or "")
    session_id = uuid.uuid4().hex
    session = {
        "id": session_id,
        "sourcePath": str(source_path),
        "sourceName": source_name,
        "sourceKind": source_kind,
        "sourceSuffix": source_path.suffix.lower() or ".docx",
        "platformFileId": platform_file_id,
        "paragraphs": paragraphs,
        "issues": [],
        "status": "ready",
        "warnings": warnings,
        "meta": {},
    }
    proofread_sessions[session_id] = session
    persist_session(session)
    return public_session(session)


def public_session(session: dict[str, Any], include_issues: bool = True) -> dict[str, Any]:
    payload = {
        "success": True,
        "sessionId": session["id"],
        "sourceName": session["sourceName"],
        "sourceKind": session["sourceKind"],
        "sourceSuffix": session["sourceSuffix"],
        "platformFileId": session.get("platformFileId"),
        "status": session.get("status") or "ready",
        "warnings": session.get("warnings") or [],
        "paragraphs": session.get("paragraphs") or [],
        "meta": session.get("meta") or {},
        "hasContent": any(str(p.get("text") or "").strip() for p in (session.get("paragraphs") or [])),
    }
    if include_issues:
        payload["issues"] = session.get("issues") or []
    return payload


def get_session(session_id: str) -> dict[str, Any]:
    session = proofread_sessions.get(session_id) or _load_session_from_disk(session_id)
    if not session:
        raise KeyError("纠错会话不存在或已过期，请重新上传文件")
    return session


def update_session_paragraphs(session: dict[str, Any], paragraphs: list[dict[str, Any]]) -> None:
    if not isinstance(paragraphs, list):
        return
    existing = {}
    for item in (session.get("paragraphs") or []):
        if not isinstance(item, dict):
            continue
        try:
            existing[int(item.get("index"))] = item
        except (TypeError, ValueError):
            continue
    cleaned: list[dict[str, Any]] = []
    for item in paragraphs:
        if not isinstance(item, dict):
            continue
        try:
            index = int(item.get("index"))
        except (TypeError, ValueError):
            continue
        prev = existing.get(index) or {}
        cleaned.append({
            "index": index,
            "text": str(item.get("text") if item.get("text") is not None else ""),
            "align": str(item.get("align") or prev.get("align") or "left"),
            "empty": not bool(str(item.get("text") or "").strip()),
            "source": str(item.get("source") or prev.get("source") or "paragraph"),
            "originalText": str(
                item.get("originalText")
                if item.get("originalText") is not None
                else prev.get("originalText")
                or prev.get("text")
                or ""
            ),
        })
    if cleaned:
        session["paragraphs"] = cleaned
        persist_session(session)


def export_session(session: dict[str, Any], paragraphs: list[dict[str, Any]] | None = None) -> tuple[Path, str, str]:
    if paragraphs is not None:
        update_session_paragraphs(session, paragraphs)
    paras = session.get("paragraphs") or []
    source_path = Path(session["sourcePath"])
    suffix = session.get("sourceSuffix") or source_path.suffix.lower() or ".docx"
    stem = Path(session.get("sourceName") or "公文").stem
    out_id = uuid.uuid4().hex
    if suffix == ".docx" and source_path.exists():
        dest = TEMP_DIR / f"{out_id}_proofread.docx"
        apply_paragraphs_to_docx(source_path, paras, dest)
        filename = f"{stem}_纠错.docx"
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        return dest, filename, media
    if suffix == ".txt":
        dest = TEMP_DIR / f"{out_id}_proofread.txt"
        export_as_txt(paras, dest)
        filename = f"{stem}_纠错.txt"
        return dest, filename, "text/plain; charset=utf-8"
    # 其余格式（pdf/图片等）优先回落为 Word，保证可编辑下载
    dest = TEMP_DIR / f"{out_id}_proofread.docx"
    export_as_docx_from_paragraphs(paras, dest, title=stem)
    filename = f"{stem}_纠错.docx"
    media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return dest, filename, media
