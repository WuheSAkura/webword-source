"""Full-chain QA: template library, API key, and per-template generation smoke tests."""
from __future__ import annotations

import json
import os
import sys
import urllib.error
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from docx import Document  # noqa: E402

from ai_writer import (  # noqa: E402
    build_upload_temp_path,
    display_name_from_temp_path,
    get_model_config,
    load_template_references,
    scan_template_library,
    sync_template_library,
)
from config import get_template_catalog  # noqa: E402

TEMP_DIR = BACKEND / "temp"
TEST_MATERIAL = TEMP_DIR / "qa_fixed_test_material.docx"
REPORT_PATH = ROOT / ".codex-runlogs" / "qa-batch-report.json"
BASE_URL = os.getenv("QA_BASE_URL", "http://127.0.0.1:8010")


def ensure_test_material() -> Path:
    TEMP_DIR.mkdir(parents=True, exist_ok=True)
    if TEST_MATERIAL.exists() and TEST_MATERIAL.stat().st_size > 0:
        return TEST_MATERIAL
    doc = Document()
    doc.add_paragraph("关于开展2024年度食品安全专项检查工作的通知")
    doc.add_paragraph("各分局、派出所：")
    doc.add_paragraph(
        "为落实上级部署，决定自2024年3月起开展专项检查。"
        "重点检查校园周边、农贸市场等重点区域，要求各单位于3月15日前报送工作方案。"
    )
    doc.add_paragraph("特此通知。")
    doc.add_paragraph("市公安局环食药侦总队")
    doc.add_paragraph("2024年3月1日")
    doc.save(str(TEST_MATERIAL))
    return TEST_MATERIAL


def validate_api_key() -> dict:
    cfg = get_model_config()
    if not cfg.get("apiKeyConfigured"):
        return {"ok": False, "reason": "DEEPSEEK_API_KEY 未配置"}
    url = str(cfg.get("requestUrl") or "").rstrip("/") + "/v1/chat/completions"
    payload = {
        "model": cfg.get("modelName"),
        "messages": [{"role": "user", "content": "回复 OK"}],
        "max_tokens": 8,
        "stream": False,
    }
    req = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {cfg.get('apiKey') or os.getenv('DEEPSEEK_API_KEY', '')}",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=45) as resp:
            body = json.loads(resp.read().decode("utf-8"))
        content = body["choices"][0]["message"]["content"]
        return {"ok": True, "model": body.get("model"), "sample": str(content)[:80]}
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        return {"ok": False, "reason": f"HTTP {exc.code}: {detail[:300]}"}
    except Exception as exc:
        return {"ok": False, "reason": str(exc)}


def validate_templates() -> dict:
    templates = sync_template_library()
    catalog = {item["id"]: item for item in get_template_catalog()}
    rows = []
    missing_refs = []
    unreadable_only = []
    for template in templates:
        refs = load_template_references(template["id"], template["sourceDir"])
        ready = [item for item in refs if item["status"] == "ready" and item["text"].strip()]
        row = {
            "id": template["id"],
            "label": template.get("label") or template["id"],
            "fileCount": template.get("fileCount", 0),
            "readyReferences": len(ready),
            "sourceDirExists": Path(template["sourceDir"]).exists(),
        }
        rows.append(row)
        if not row["sourceDirExists"]:
            missing_refs.append(template["id"])
        if row["fileCount"] > 0 and not ready:
            unreadable_only.append(template["id"])
    return {
        "count": len(templates),
        "expected": 15,
        "ok": len(templates) == 15 and not missing_refs,
        "missingSourceDirs": missing_refs,
        "unreadableOnly": unreadable_only,
        "templates": rows,
    }


def post_generate(template_id: str, material: Path, long_name: bool = False) -> dict:
    cfg = get_model_config()
    boundary = "----qaBoundary7MA4YWxk"
    filename = "④被告人韦忠南犯故意伤害罪测试材料.docx" if long_name else material.name
    file_bytes = material.read_bytes()
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="prompt"\r\n\r\n'
        f"请根据材料撰写一份正式公文，按模板结构组织内容，不得照搬原文。\r\n"
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="template_id"\r\n\r\n'
        f"{template_id}\r\n"
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="speed_mode"\r\n\r\n'
        f"fast\r\n"
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="files"; filename="{filename}"\r\n'
        f"Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n"
    ).encode("utf-8") + file_bytes + f"\r\n--{boundary}--\r\n".encode("utf-8")
    req = urllib.request.Request(
        f"{BASE_URL}/api/ai/generate",
        data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=180) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
        preview = str(payload.get("preview") or "")
        validation = (
            ((payload.get("documents") or [{}])[0].get("generation") or {}).get("validation")
            or {}
        )
        return {
            "ok": True,
            "documentId": payload.get("documentId"),
            "previewChars": len(preview),
            "sourceCopyCoverage": validation.get("sourceCopyCoverage"),
            "copiedParagraphRatio": validation.get("copiedParagraphRatio"),
            "structureCoverage": validation.get("structureCoverage"),
        }
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        try:
            message = json.loads(detail).get("detail", detail)
        except json.JSONDecodeError:
            message = detail
        return {"ok": False, "reason": str(message)[:500]}
    except Exception as exc:
        return {"ok": False, "reason": str(exc)}


def validate_upload_path_encoding() -> dict:
    upload_path = build_upload_temp_path(TEMP_DIR, "④被告人韦忠南犯故意伤害罪测试材料.docx", ".docx")
    upload_path.write_bytes(b"demo")
    display = display_name_from_temp_path(upload_path)
    meta_exists = upload_path.with_suffix(upload_path.suffix + ".upload.json").exists()
    ok = upload_path.exists() and meta_exists and display.endswith(".docx")
    upload_path.unlink(missing_ok=True)
    upload_path.with_suffix(upload_path.suffix + ".upload.json").unlink(missing_ok=True)
    return {"ok": ok, "displayName": display, "pathLen": len(str(upload_path))}


def main() -> int:
    material = ensure_test_material()
    report = {
        "uploadPath": validate_upload_path_encoding(),
        "templates": validate_templates(),
        "apiKey": validate_api_key(),
        "generation": [],
        "longFilenameSmoke": None,
    }

    if not report["apiKey"]["ok"]:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 2

    for template in report["templates"]["templates"]:
        template_id = template["id"]
        if template["readyReferences"] < 1:
            report["generation"].append({
                "templateId": template_id,
                "label": template["label"],
                "ok": False,
                "skipped": True,
                "reason": "无可解析范文（多为 wps/ofd，需转 DOCX/PDF）",
            })
            continue
        result = post_generate(template_id, material)
        result["templateId"] = template_id
        result["label"] = template["label"]
        report["generation"].append(result)
        print(f"[{template_id}] {'OK' if result.get('ok') else 'FAIL'}", flush=True)

    notice = next((item for item in report["templates"]["templates"] if item["id"] == "notice"), None)
    if notice and notice["readyReferences"] > 0:
        report["longFilenameSmoke"] = post_generate("notice", material, long_name=True)

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    ok_count = sum(1 for item in report["generation"] if item.get("ok"))
    skip_count = sum(1 for item in report["generation"] if item.get("skipped"))
    fail_count = len(report["generation"]) - ok_count - skip_count
    print(
        f"QA done: ok={ok_count} skipped={skip_count} fail={fail_count} report={REPORT_PATH}",
        flush=True,
    )
    return 0 if fail_count == 0 and report["templates"]["ok"] and report["apiKey"]["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
