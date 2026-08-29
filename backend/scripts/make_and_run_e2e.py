"""5 文件 × 5 文种，不勾选降级，跑通 AI 成文→平台整理全流程。"""
from __future__ import annotations

import json
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent / "temp" / "e2e_materials"
BASE = "http://127.0.0.1:8010"

CASES = [
    {
        "file": "case1_专项检查.txt.docx",
        "template_id": "letter",
        "title": "关于商请协助开展专项检查的情况材料",
        "body": (
            "根据上级统一部署，我单位拟于近期开展食品安全专项检查。\n"
            "一是检查时间安排在3月中旬；二是重点核查台账资料与现场落实情况；"
            "三是请各相关单位做好迎检准备，并于检查结束后5个工作日内反馈有关情况。\n"
            "落款单位：某市市场监管局\n成文日期：2024年3月1日"
        ),
        "prompt": "请据此创作一份正式函。",
    },
    {
        "file": "case2_工作推进.txt.docx",
        "template_id": "notice",
        "title": "关于推进年度重点工作的情况说明",
        "body": (
            "为落实年度目标任务，现将有关工作安排明确如下。\n"
            "一、总体要求：围绕主责主业，压实责任。\n"
            "二、重点任务：开展隐患排查、完善台账、组织培训。\n"
            "三、时间安排：4月底前完成第一阶段，6月底前形成阶段性报告。\n"
            "请各处室认真落实。"
        ),
        "prompt": "请据此创作一份通知。",
    },
    {
        "file": "case3_工作总结.txt.docx",
        "template_id": "report",
        "title": "2024年度重点工作推进情况",
        "body": (
            "本年度围绕年度目标任务扎实推进各项工作。\n"
            "一、工作成效：完成专项检查12次，整改问题35项，培训干部80人次。\n"
            "二、存在问题：基础台账不够规范，协同机制有待完善。\n"
            "三、下步安排：补齐短板，强化督导，确保任务闭环。"
        ),
        "prompt": "请据此创作一份报告。",
    },
    {
        "file": "case4_实施意见.txt.docx",
        "template_id": "opinion",
        "title": "关于加强行业监管工作的意见材料",
        "body": (
            "为进一步规范行业监管，现提出如下意见。\n"
            "一、压实主体责任，明确属地管理要求。\n"
            "二、健全风险预警机制，做到早发现早处置。\n"
            "三、强化部门协同，形成监管合力。\n"
            "请结合实际认真贯彻执行。"
        ),
        "prompt": "请据此创作一份意见。",
    },
    {
        "file": "case5_张某某周某某案情.txt.docx",
        "template_id": "minutes",
        "title": "张某某、周某某妨害公务案情况摘录",
        "body": (
            "被告人张某某、周某某因妨害公务一案，经法院审理查明：二人在执法现场阻碍公务人员依法执行职务。\n"
            "法院依法作出判决。现就有关情况通报如下，请各单位引以为戒，配合执法。"
        ),
        "prompt": "请据此创作一份纪要式情况材料，保留材料中的脱敏姓名。",
    },
]


def write_docx(path: Path, title: str, body: str) -> None:
    doc = Document()
    doc.add_paragraph(title)
    for para in body.split("\n"):
        doc.add_paragraph(para)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def http_json(method: str, url: str, timeout: int = 30) -> dict:
    req = urllib.request.Request(url, method=method)
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def post_multipart(url: str, fields: dict[str, str], files: list[tuple[str, Path]], timeout: int = 900) -> dict:
    boundary = "----WebKitFormBoundaryE2E7Gongwen"
    body = bytearray()
    for key, value in fields.items():
        body.extend(f"--{boundary}\r\n".encode())
        body.extend(f'Content-Disposition: form-data; name="{key}"\r\n\r\n'.encode())
        body.extend(f"{value}\r\n".encode("utf-8"))
    for field_name, path in files:
        data = path.read_bytes()
        body.extend(f"--{boundary}\r\n".encode())
        body.extend(
            f'Content-Disposition: form-data; name="{field_name}"; filename="{path.name}"\r\n'.encode("utf-8")
        )
        body.extend(b"Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n")
        body.extend(data)
        body.extend(b"\r\n")
    body.extend(f"--{boundary}--\r\n".encode())
    req = urllib.request.Request(
        url,
        data=bytes(body),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        try:
            return json.loads(detail)
        except Exception:
            return {"success": False, "error": f"HTTP {exc.code}: {detail[:800]}"}


def main() -> int:
    health = http_json("GET", f"{BASE}/api/health")
    print("health:", health)
    model = http_json("GET", f"{BASE}/api/ai/model-config")
    print("modelConfigured:", model.get("apiKeyConfigured"), "model:", model.get("modelName"))
    if not model.get("apiKeyConfigured"):
        print("FAIL: model API key not configured")
        return 2

    results = []
    for index, case in enumerate(CASES, start=1):
        path = ROOT / case["file"]
        write_docx(path, case["title"], case["body"])
        print(f"\n=== [{index}/5] {case['file']} -> template={case['template_id']} ===")
        started = time.time()
        payload = post_multipart(
            f"{BASE}/api/ai/generate",
            {
                "prompt": case["prompt"],
                "template_id": case["template_id"],
                "speed_mode": "fast",
                "temperature": "0.3",
                "strict_reference_isolation": "false",
                "allow_degradation": "false",
            },
            [("material_files", path)],
            timeout=900,
        )
        elapsed = round(time.time() - started, 1)
        ok = bool(payload.get("success"))
        docs = payload.get("documents") or ([] if not payload.get("fileId") else [payload])
        preview_len = 0
        warnings = []
        if docs:
            preview_len = len(str(docs[0].get("preview") or ""))
            warnings = docs[0].get("warnings") or payload.get("warnings") or []
        item = {
            "case": case["file"],
            "template_id": case["template_id"],
            "ok": ok,
            "elapsedSec": elapsed,
            "documentCount": len(docs),
            "previewChars": preview_len,
            "fileId": (docs[0].get("fileId") if docs else None) or payload.get("fileId"),
            "error": None if ok else (payload.get("error") or payload.get("detail") or str(payload)[:500]),
            "warningCount": len(warnings),
            "warningsHead": warnings[:4],
        }
        results.append(item)
        print(json.dumps(item, ensure_ascii=False, indent=2))

    out = ROOT / "e2e_results.json"
    out.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    passed = sum(1 for item in results if item["ok"])
    print(f"\nSUMMARY: {passed}/5 passed")
    print("saved:", out)
    return 0 if passed == 5 else 1


if __name__ == "__main__":
    sys.exit(main())
