"""单次快路径耗时验证：不勾选降级，期望仅 1 次成文模型调用。"""
from __future__ import annotations

import json
import time
import urllib.error
import urllib.request
from pathlib import Path

from docx import Document

BASE = "http://127.0.0.1:8010"
OUT = Path(__file__).resolve().parent.parent / "temp" / "e2e_materials" / "speed_case.docx"


def main() -> int:
    doc = Document()
    doc.add_paragraph("关于推进专项检查工作的情况材料")
    doc.add_paragraph("根据工作安排，拟于3月开展专项检查。一是核查台账；二是现场抽查；三是汇总反馈。")
    doc.add_paragraph("请相关单位配合，并于月底前报送情况。")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)

    boundary = "----WebKitFormBoundarySpeed"
    body = bytearray()
    fields = {
        "prompt": "请据此创作一份正式函。",
        "template_id": "letter",
        "speed_mode": "fast",
        "temperature": "0.3",
        "strict_reference_isolation": "false",
        "allow_degradation": "false",
    }
    for key, value in fields.items():
        body.extend(f"--{boundary}\r\n".encode())
        body.extend(f'Content-Disposition: form-data; name="{key}"\r\n\r\n'.encode())
        body.extend(f"{value}\r\n".encode("utf-8"))
    data = OUT.read_bytes()
    body.extend(f"--{boundary}\r\n".encode())
    body.extend(
        f'Content-Disposition: form-data; name="material_files"; filename="{OUT.name}"\r\n'.encode("utf-8")
    )
    body.extend(b"Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n")
    body.extend(data)
    body.extend(b"\r\n")
    body.extend(f"--{boundary}--\r\n".encode())
    req = urllib.request.Request(
        f"{BASE}/api/ai/generate",
        data=bytes(body),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    started = time.time()
    try:
        with urllib.request.urlopen(req, timeout=600) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        payload = json.loads(exc.read().decode("utf-8", errors="ignore") or "{}")
        payload.setdefault("success", False)
        payload.setdefault("error", f"HTTP {exc.code}")
    elapsed = round(time.time() - started, 1)
    docs = payload.get("documents") or ([] if not payload.get("fileId") else [payload])
    gen = (docs[0].get("generation") if docs else {}) or payload.get("generation") or {}
    stages = (docs[0].get("stages") if docs else {}) or payload.get("stages") or {}
    result = {
        "ok": bool(payload.get("success")),
        "elapsedSec": elapsed,
        "prepSeconds": stages.get("prepSeconds"),
        "planSeconds": stages.get("planSeconds") or gen.get("planSeconds"),
        "draftingSeconds": stages.get("draftingSeconds"),
        "rewriteRetries": gen.get("rewriteRetries"),
        "similarityRewrites": gen.get("similarityRewrites"),
        "draftAttempts": len((gen.get("attempts") or [])),
        "plannerFallback": (gen.get("planner") or {}).get("fallback"),
        "plannerReason": (gen.get("planner") or {}).get("fallbackReason"),
        "plannerFinish": (gen.get("planner") or {}).get("finishReason"),
        "copiedParagraphRatio": ((gen.get("validation") or {}).get("copiedParagraphRatio")),
        "sourceCopyCoverage": ((gen.get("validation") or {}).get("sourceCopyCoverage")),
        "previewHead": ((docs[0].get("preview") if docs else "") or "")[:280],
        "fileId": (docs[0].get("fileId") if docs else None) or payload.get("fileId"),
        "error": None if payload.get("success") else (payload.get("error") or payload.get("detail")),
        "warningsHead": ((docs[0].get("warnings") if docs else []) or [])[:8],
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
