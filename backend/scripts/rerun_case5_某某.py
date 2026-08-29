"""单独复跑含「某某」脱敏名的第5例，验证占位符不再硬拦。"""
from __future__ import annotations

import json
import time
import urllib.error
import urllib.request
from pathlib import Path

from docx import Document

BASE = "http://127.0.0.1:8010"
ROOT = Path(__file__).resolve().parent.parent / "temp" / "e2e_materials"
PATH = ROOT / "case5_张某某周某某案情.txt.docx"


def write_docx() -> None:
    doc = Document()
    doc.add_paragraph("张某某、周某某妨害公务案情况摘录")
    doc.add_paragraph(
        "被告人张某某、周某某因妨害公务一案，经法院审理查明：二人在执法现场阻碍公务人员依法执行职务。"
    )
    doc.add_paragraph("法院依法作出判决。现就有关情况通报如下，请各单位引以为戒，配合执法。")
    PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(PATH)


def post() -> dict:
    boundary = "----WebKitFormBoundaryCase5"
    body = bytearray()
    fields = {
        "prompt": "请据此创作一份纪要式情况材料，保留材料中的脱敏姓名。",
        "template_id": "minutes",
        "speed_mode": "fast",
        "temperature": "0.3",
        "strict_reference_isolation": "false",
        "allow_degradation": "false",
    }
    for key, value in fields.items():
        body.extend(f"--{boundary}\r\n".encode())
        body.extend(f'Content-Disposition: form-data; name="{key}"\r\n\r\n'.encode())
        body.extend(f"{value}\r\n".encode("utf-8"))
    data = PATH.read_bytes()
    body.extend(f"--{boundary}\r\n".encode())
    body.extend(
        f'Content-Disposition: form-data; name="material_files"; filename="{PATH.name}"\r\n'.encode("utf-8")
    )
    body.extend(b"Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n")
    body.extend(data)
    body.extend(b"\r\n")
    body.extend(f"--{boundary}--\r\n".encode())
    req = urllib.request.Request(
        f"{BASE}/api/ai/generate",
        data=bytes(body),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=900) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        try:
            return json.loads(detail)
        except Exception:
            return {"success": False, "error": detail[:800]}


def main() -> int:
    write_docx()
    started = time.time()
    payload = post()
    elapsed = round(time.time() - started, 1)
    docs = payload.get("documents") or ([] if not payload.get("fileId") else [payload])
    preview = str((docs[0].get("preview") if docs else "") or payload.get("preview") or "")
    ok = bool(payload.get("success"))
    result = {
        "ok": ok,
        "elapsedSec": elapsed,
        "template_id": "minutes",
        "previewChars": len(preview),
        "has某某": ("某某" in preview),
        "fileId": (docs[0].get("fileId") if docs else None) or payload.get("fileId"),
        "error": None if ok else (payload.get("error") or payload.get("detail") or str(payload)[:500]),
        "warningsHead": ((docs[0].get("warnings") if docs else []) or [])[:5],
        "previewHead": preview[:300],
    }
    out = ROOT / "e2e_case5_result.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
